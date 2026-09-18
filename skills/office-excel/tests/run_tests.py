#!/usr/bin/env python3
"""office-excel çevrimdışı testleri (stdlib unittest). Örnek dosyalar test sırasında geçici klasörde üretilir.

    python skills/office-excel/tests/run_tests.py

openpyxl kuruluysa bağımsız bir üreticiyle çapraz kontrol de koşar (bizim yazdığımızı openpyxl okur, openpyxl'in
yazdığını biz okuruz); resim testi için ayrıca Pillow gerekir. Eksik olan testler SKIP olarak raporlanır.
Çıkış: 0 tümü geçti · 1 en az bir başarısız.
"""
from __future__ import annotations

import datetime as dt
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
import office_excel as OE  # noqa: E402
import xlsx_lite as X  # noqa: E402

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
HAS_PIL = importlib.util.find_spec("PIL") is not None


def cli(*args, python_code: str | None = None):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
    cmd = [sys.executable, "-c", python_code] if python_code else [sys.executable, str(SCRIPTS / "office_excel.py")]
    p = subprocess.run(cmd + [str(a) for a in args], capture_output=True, text=True, encoding="utf-8",
                       stdin=subprocess.DEVNULL, env=env)
    return p.returncode, p.stdout, p.stderr


class Tmp(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()


class XlsxLiteTests(Tmp):
    ROWS = [["Malzeme", "Açıklama", "Miktar", "Aktif", "Tarih"],
            ["000010", "Çelik İğne şğü", 12.5, True, "2026-09-13"],
            ["000020", None, 3, False, None]]

    def test_col_letters_roundtrip(self):
        for i, s in ((0, "A"), (25, "Z"), (26, "AA"), (701, "ZZ"), (702, "AAA")):
            self.assertEqual(X.col_letters(i), s)
            self.assertEqual(X.col_index(s), i)

    def test_write_read_roundtrip_and_sheet_name(self):
        path = self.d / "a.xlsx"
        warnings = X.write_xlsx(path, [("Stok/2026", self.ROWS)])
        self.assertTrue(any("Stok_2026" in w for w in warnings))
        name, rows = X.read_rows(path)
        self.assertEqual(name, "Stok_2026")
        self.assertEqual(rows[0], self.ROWS[0])
        self.assertEqual(rows[1], ["000010", "Çelik İğne şğü", 12.5, True, "2026-09-13"])
        self.assertEqual(rows[2], ["000020", None, 3, False])  # sondaki boş hücre okunmaz

    def test_bad_zip(self):
        bad = self.d / "eski.xlsx"
        bad.write_bytes(b"not a zip")
        with self.assertRaises(X.XlsxError):
            X.read_rows(bad)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl yok")
    def test_openpyxl_reads_our_file(self):
        import openpyxl
        path = self.d / "a.xlsx"
        X.write_xlsx(path, [("Stok", self.ROWS)])
        ws = openpyxl.load_workbook(path).active
        self.assertEqual(ws["B2"].value, "Çelik İğne şğü")
        self.assertEqual(ws["C2"].value, 12.5)
        self.assertIs(ws["D3"].value, False)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:E3")
        self.assertTrue(ws["A1"].font.b)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl yok")
    def test_we_read_openpyxl_file(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Veri"
        ws.append(["Tarih", "Zaman", "Metin", "Tutar"])
        ws.append([dt.date(2026, 9, 13), dt.datetime(2026, 9, 13, 14, 30), "İzmir", 1234.5])
        wb.create_sheet("İkinci").append(["x"])
        path = self.d / "o.xlsx"
        wb.save(path)
        self.assertEqual(X.list_sheets(path), ["Veri", "İkinci"])
        _name, rows = X.read_rows(path, "Veri")
        self.assertEqual(rows[1], ["2026-09-13", "2026-09-13 14:30:00", "İzmir", 1234.5])
        self.assertEqual(X.read_rows(path, "2")[0], "İkinci")


class TableAndProfileTests(Tmp):
    def test_csv_cp1254_semicolon(self):
        p = self.d / "tr.csv"
        p.write_bytes("Malzeme;Tutar\n000010;1,5\nŞiş;2\n".encode("cp1254"))
        t = OE.load_table(p)
        self.assertIn("cp1254", t.notes[0])
        self.assertIn("';'", t.notes[0])
        self.assertEqual(t.rows[1][0], "Şiş")
        self.assertEqual(OE.to_number(t.rows[0][1]), 1.5)

    def test_header_row_warning_and_fix(self):
        path = self.d / "bex.xlsx"
        grid = [["Filtre: 2026"], [], ["Bölge", "Ciro", "Adet", "Pay"], ["Ege", 10, 1, 0.5], ["Marmara", 20, 2, 0.5]]
        X.write_xlsx(path, [("Rapor", grid)], header=False)
        rc, out, _ = cli("profile", path, "--json")
        self.assertEqual(rc, 0)
        self.assertTrue(any("--header-row" in w for w in json.loads(out)["uyarılar"]))
        rc, out, _ = cli("profile", path, "--json", "--header-row", "3")
        data = json.loads(out)
        self.assertEqual(rc, 0)
        self.assertEqual([c["kolon"] for c in data["kolonlar"]], ["Bölge", "Ciro", "Adet", "Pay"])
        self.assertFalse(any("--header-row" in w for w in data["uyarılar"]))
        ciro = data["kolonlar"][1]
        self.assertEqual((ciro["tip"], ciro["min"], ciro["max"], ciro["toplam"]), ("sayı", 10, 20, 30))

    def test_profile_no_samples_hides_values(self):
        p = self.d / "k.csv"
        p.write_text("Ad,TCKN\nAyşe,10000000146\n", encoding="utf-8")
        rc, out, _ = cli("profile", p, "--json", "--no-samples")
        self.assertEqual(rc, 0)
        self.assertNotIn("10000000146", out)
        self.assertNotIn("Ayşe", out)


class TransformTests(Tmp):
    def setUp(self):
        super().setUp()
        self.src = self.d / "satis.csv"
        self.src.write_text("Bölge,Ürün,Tutar\nİSTANBUL,Vida,150\nİSTANBUL,Somun,50\nEge,Vida,300\nEge,Pul,120\n"
                            "Ege,Pul,120\n", encoding="utf-8")

    def test_pipeline(self):
        out = self.d / "sonuc.json"
        rc, stdout, err = cli("transform", self.src, out, "--where", "Tutar|gt|100", "--dedupe",
                              "--group-by", "Bölge", "--agg", "Toplam=sum:Tutar", "--agg", "Adet=count:Ürün",
                              "--sort", "Toplam:desc", "--rename", "Bölge=Bolge")
        self.assertEqual(rc, 0, err)
        data = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(data, [{"Bolge": "Ege", "Toplam": 420, "Adet": 2}, {"Bolge": "İSTANBUL", "Toplam": 150, "Adet": 1}])
        self.assertIn("satır 5 → 2", stdout)

    def test_turkish_contains_and_keep(self):
        out = self.d / "ist.csv"
        rc, _, err = cli("transform", self.src, out, "--where", "Bölge|contains|istanbul", "--keep", "Ürün")
        self.assertEqual(rc, 0, err)
        self.assertEqual(out.read_text(encoding="utf-8-sig").split(), ["Ürün", "Vida", "Somun"])

    def test_input_never_overwritten_and_force(self):
        rc, _, err = cli("transform", self.src, self.src, "--dedupe")
        self.assertEqual(rc, 3)
        self.assertIn("girdi dosyasının üzerine yazılmaz", err)
        out = self.d / "x.xlsx"
        self.assertEqual(cli("convert", self.src, out)[0], 0)
        self.assertEqual(cli("convert", self.src, out)[0], 3)
        self.assertEqual(cli("convert", self.src, out, "--force")[0], 0)

    def test_bad_where(self):
        rc, _, err = cli("transform", self.src, self.d / "y.csv", "--where", "Tutar > 100")
        self.assertEqual(rc, 3)
        self.assertIn("Kolon|op|değer", err)


class CompareTests(Tmp):
    def setUp(self):
        super().setUp()
        self.old = self.d / "eski.csv"
        self.old.write_text("Malzeme;Ad;Fiyat\n000010;Vida;1.5\n000020;Somun;2\n000030;Pul;0.5\n", encoding="utf-8")
        self.new = self.d / "yeni.xlsx"
        X.write_xlsx(self.new, [("Liste", [["Malzeme", "Ad", "Fiyat", "Yeni"], [10, "Vida", 1.5, "x"],
                                           ["000020", "Somun", 2.5, None], ["000040", "Cıvata", 3, None],
                                           ["000040", "tekrar", 1, None]])])

    def summary(self, stdout):
        return dict(line.split(": ", 1) for line in stdout.splitlines() if ": " in line and not line.startswith(" "))

    def test_trim_mode_leading_zeros_differ(self):
        rc, out, err = cli("compare", self.old, self.new, "--key", "Malzeme")
        self.assertEqual(rc, 0, err)
        s = self.summary(out)
        self.assertEqual((s["eklenen"], s["silinen"], s["değişen_satır"], s["tekrarlanan_anahtar_yeni"]), ("2", "2", "1", "1"))

    def test_number_mode_and_report(self):
        rep = self.d / "fark.xlsx"
        rc, out, err = cli("compare", self.old, self.new, "--key", "Malzeme", "--key-mode", "number", "--output", rep)
        self.assertEqual(rc, 0, err)
        s = self.summary(out)
        self.assertEqual((s["eklenen"], s["silinen"], s["değişen_satır"], s["aynı_satır"]), ("1", "1", "1", "1"))
        self.assertEqual(s["yalnız_yenide_kolon"], "['Yeni']")
        self.assertIn("'2' → 2.5", out)
        self.assertEqual(X.list_sheets(rep), ["Ozet", "Eklenen", "Silinen", "Degisen", "Tekrarlanan_Anahtar"])
        self.assertEqual(X.read_rows(rep, "Degisen")[1], [["Malzeme", "Kolon", "Eski", "Yeni"], ["20", "Fiyat", "2", 2.5]])

    def test_summary_only_writes_nothing(self):
        rep = self.d / "yok.xlsx"
        rc, out, _ = cli("compare", self.old, self.new, "--key", "Malzeme", "--summary-only", "--output", rep)
        self.assertEqual(rc, 0)
        self.assertFalse(rep.exists())
        self.assertNotIn("değişti", out)


def _fixture_with_drawing(path: Path) -> None:
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
    rt = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    pic = ('<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{i}" name="P{i}"/><xdr:cNvPicPr/></xdr:nvPicPr>'
           '<xdr:blipFill><a:blip r:embed="{rid}"/></xdr:blipFill><xdr:spPr/></xdr:pic>')
    drawing = ('<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="%s">'
               '<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row>'
               '<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row>'
               '<xdr:rowOff>0</xdr:rowOff></xdr:to>%s<xdr:clientData/></xdr:twoCellAnchor>'
               '<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row>'
               '<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="1" cy="1"/>%s<xdr:clientData/></xdr:oneCellAnchor>'
               '</xdr:wsDr>') % (ns_r, pic.format(i=1, rid="rId1"), pic.format(i=2, rid="rId2"))
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("_rels/.rels", f'<Relationships xmlns="{ns_pkg}"><Relationship Id="rId1" Type="{rt}officeDocument" Target="xl/workbook.xml"/></Relationships>')
        zf.writestr("xl/workbook.xml", f'<workbook xmlns="{ns_main}" xmlns:r="{ns_r}"><sheets><sheet name="Ürünler" sheetId="1" r:id="rId1"/></sheets></workbook>')
        zf.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{ns_pkg}"><Relationship Id="rId1" Type="{rt}worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        zf.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{ns_main}" xmlns:r="{ns_r}"><sheetData/><drawing r:id="rId1"/></worksheet>')
        zf.writestr("xl/worksheets/_rels/sheet1.xml.rels", f'<Relationships xmlns="{ns_pkg}"><Relationship Id="rId1" Type="{rt}drawing" Target="../drawings/drawing1.xml"/></Relationships>')
        zf.writestr("xl/drawings/drawing1.xml", drawing)
        zf.writestr("xl/drawings/_rels/drawing1.xml.rels", f'<Relationships xmlns="{ns_pkg}"><Relationship Id="rId1" Type="{rt}image" Target="../media/image1.png"/><Relationship Id="rId2" Type="{rt}image" Target="../media/image1.png"/></Relationships>')
        zf.writestr("xl/media/image1.png", b"\x89PNG\r\n\x1a\nAAAA")
        zf.writestr("xl/media/image2.png", b"\x89PNG\r\n\x1a\nBBBB")


class ImagesTests(Tmp):
    def test_anchors_dedupe_and_unplaced(self):
        src = self.d / "fiyat.xlsx"
        _fixture_with_drawing(src)
        rc, out, err = cli("images", src, "--json")
        self.assertEqual(rc, 0, err)
        data = json.loads(out)
        self.assertEqual(data["benzersiz_dosya"], 2)
        cells = {(r["sayfa"], r["hücre"]): r["dosya"] for r in data["resimler"]}
        self.assertEqual(cells[("Ürünler", "B3")], "Ürünler_B3_1.png")
        self.assertEqual(cells[("Ürünler", "A1")], "Ürünler_B3_1.png")  # aynı içerik → tek dosya
        self.assertEqual(cells[(None, None)], "konumsuz_image2_1.png")
        self.assertEqual(sorted(p.name for p in (self.d / "fiyat_images").iterdir()), ["konumsuz_image2_1.png", "Ürünler_B3_1.png"])
        self.assertEqual(cli("images", src)[0], 3)  # dolu klasör --force olmadan ezilmez

    def test_no_media(self):
        src = self.d / "bos.xlsx"
        X.write_xlsx(src, [("S", [["a"], [1]])])
        rc, out, _ = cli("images", src)
        self.assertEqual(rc, 0)
        self.assertIn("gömülü resim yok", out)
        self.assertFalse((self.d / "bos_images").exists())

    @unittest.skipUnless(HAS_OPENPYXL and HAS_PIL, "openpyxl + Pillow yok")
    def test_openpyxl_made_image(self):
        import openpyxl
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image
        png = self.d / "k.png"
        Image.new("RGB", (8, 8), (200, 0, 0)).save(png)
        wb = openpyxl.Workbook()
        wb.active.add_image(XLImage(str(png)), "C5")
        src = self.d / "op.xlsx"
        wb.save(src)
        rc, out, err = cli("images", src, "--json")
        self.assertEqual(rc, 0, err)
        self.assertEqual([r["hücre"] for r in json.loads(out)["resimler"]], ["C5"])


class ReportAndCliTests(Tmp):
    def setUp(self):
        super().setUp()
        self.src = self.d / "ay.csv"
        self.src.write_text("Ay,Ciro,Maliyet\nOcak,1000.5,700\nŞubat,1200,800\n", encoding="utf-8")

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl yok")
    def test_report(self):
        import openpyxl
        out = self.d / "rapor.xlsx"
        rc, _, err = cli("report", self.src, out, "--title", "Aylık Ciro", "--number-format", "Ciro=#,##0.00",
                         "--chart", "bar", "--chart-x", "Ay", "--chart-y", "Ciro", "--chart-y", "Maliyet")
        self.assertEqual(rc, 0, err)
        ws = openpyxl.load_workbook(out).active
        self.assertEqual(ws["A1"].value, "Aylık Ciro")
        self.assertEqual(ws["B4"].value, 1000.5)
        self.assertEqual(ws["B4"].number_format, "#,##0.00")
        self.assertEqual(ws.freeze_panes, "A4")
        self.assertEqual(len(ws._charts), 1)

    def test_report_dependency_missing_exit_4(self):
        code = ("import sys, runpy; sys.modules['openpyxl'] = None; "
                f"sys.argv = ['office_excel.py', 'report', {str(self.src)!r}, {str(self.d / 'r.xlsx')!r}]; "
                f"runpy.run_path({str(SCRIPTS / 'office_excel.py')!r}, run_name='__main__')")
        rc, _, err = cli(python_code=code)
        self.assertEqual(rc, 4)
        self.assertIn("EKSİK BAĞIMLILIK", err)

    def test_usage_and_data_errors(self):
        self.assertEqual(cli()[0], 3)
        self.assertEqual(cli("profile", self.d / "yok.csv")[0], 3)
        xls = self.d / "eski.xls"
        xls.write_bytes(b"\xd0\xcf\x11\xe0")
        rc, _, err = cli("profile", xls)
        self.assertEqual(rc, 1)
        self.assertIn(".xlsx olarak kaydedip", err)


if __name__ == "__main__":
    unittest.main(verbosity=2)

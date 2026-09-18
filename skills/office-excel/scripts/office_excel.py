#!/usr/bin/env python3
"""office_excel.py — Excel / CSV / JSON tablo işleri (office-excel skill'i).

Alt komutlar:
  profile    sayfalar, kolonlar, tip, dolu/boş sayısı, örnek değerler, sayısal min/max/toplam, uyarılar
  convert    csv/tsv/json/xlsx → xlsx/csv/tsv/json (xlsx'te biçimli başlık, dondurma, filtre)
  transform  süz → tekrarları at → grupla/topla → sırala → yeniden adlandır → kolon seç/at  (bu sırayla)
  compare    iki tabloyu anahtar kolon(lar)la karşılaştırır: eklenen / silinen / değişen hücreler
  images     .xlsx içindeki gömülü resimleri sayfa + hücre adıyla çıkarır
  report     biçimli rapor: başlık, sayı biçimi, zebra satır, grafik  (isteğe bağlı bağımlılık: openpyxl)

`report` dışındaki her şey yalnız Python standart kütüphanesiyle çalışır.
Girdi dosyası asla değiştirilmez; var olan çıktı `--force` olmadan ezilmez.
Çıkış kodu: 0 başarı · 1 veri/işlem hatası · 3 kullanım hatası · 4 isteğe bağlı bağımlılık yok.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlsx_lite as X  # noqa: E402

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

EXCEL_EXT = {".xlsx", ".xlsm"}
TEXT_EXT = {".csv", ".tsv", ".txt"}
XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_NUM = re.compile(r"^[+-]?(\d+([.,]\d+)?|[.,]\d+)$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$")


class UsageError(Exception):
    pass


class DataError(Exception):
    pass


class DepMissing(Exception):
    pass


class Table:
    def __init__(self, name: str, columns: list[str], rows: list[list], notes: list[str] | None = None):
        self.name, self.columns, self.rows, self.notes = name, columns, rows, notes or []
        self.blank_headers = 0  # başlığı boş olup KolonN adı verilen kolon sayısı (_from_grid doldurur)

    def col(self, name: str) -> int:
        try:
            return self.columns.index(name)
        except ValueError:
            raise UsageError(f"kolon yok: {name!r} · mevcut kolonlar: {', '.join(self.columns)}") from None


# ── yardımcılar ───────────────────────────────────────────────────────────────────────────────────

def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def tr_lower(s: str) -> str:
    return s.replace("İ", "i").replace("I", "ı").lower()


def to_number(v):
    """Sayı ya da tek ondalık ayraçlı (nokta/virgül) sayı metni → sayı; değilse None. Binlik ayracı çözülmez."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.strip()
        if _NUM.match(s):
            f = float(s.replace(",", "."))
            return int(f) if f.is_integer() and "." not in s and "," not in s else f
    return None


def kind(v) -> str | None:
    if blank(v):
        return None
    if isinstance(v, bool):
        return "mantıksal"
    if to_number(v) is not None:
        return "sayı"
    if isinstance(v, str) and _ISO_DATE.match(v.strip()):
        return "tarih"
    return "metin"


def show(v):
    return "" if v is None else v


# ── okuma / yazma ─────────────────────────────────────────────────────────────────────────────────

def _from_grid(name: str, grid: list[list], header_row: int) -> Table:
    if header_row < 1:
        raise UsageError("--header-row 1 ya da daha büyük olmalı")
    if len(grid) < header_row:
        raise DataError(f"{name}: başlık satırı {header_row} yok (sayfada {len(grid)} satır var)")
    header = list(grid[header_row - 1])
    body = [list(r) for r in grid[header_row:]]
    width = 0
    for j in range(max([len(header)] + [len(r) for r in body])):
        in_header = j < len(header) and not blank(header[j])
        in_body = any(j < len(r) and not blank(r[j]) for r in body)
        if in_header or in_body:
            width = j + 1
    header = (header + [None] * width)[:width]
    columns, notes, seen = [], [], {}
    for j, v in enumerate(header):
        name_j = "" if v is None else str(v).strip()
        if v is not None and str(v) != name_j:
            notes.append(f"başlıktaki baş/son boşluk kırpıldı: {str(v)!r}")
        if not name_j:
            name_j = f"Kolon{j + 1}"
            notes.append(f"{X.col_letters(j)} kolonunun başlığı boş → {name_j}")
        if name_j in seen:
            seen[name_j] += 1
            renamed = f"{name_j}_{seen[name_j]}"
            notes.append(f"tekrarlanan başlık {name_j!r} → {renamed}")
            name_j = renamed
        else:
            seen[name_j] = 1
        columns.append(name_j)
    rows, empty = [], 0
    for r in body:
        r = (r + [None] * width)[:width]
        if all(blank(v) for v in r):
            empty += 1
            continue
        rows.append(r)
    if empty:
        notes.append(f"{empty} tamamen boş satır atlandı")
    table = Table(name, columns, rows, notes)
    table.blank_headers = sum(1 for v in header if blank(v))
    return table


def _read_text(path: Path) -> tuple[str, str]:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1254"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def _sniff(text: str, ext: str) -> str:
    if ext == ".tsv":
        return "\t"
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample else ""
        return max(",;\t|", key=first.count) if first else ","


def load_table(path, sheet=None, header_row: int = 1, json_key: str | None = None) -> Table:
    p = Path(path)
    if not p.is_file():
        raise UsageError(f"dosya yok: {p}")
    ext = p.suffix.lower()
    if ext in EXCEL_EXT:
        try:
            name, grid = X.read_rows(p, sheet)
        except X.XlsxError as exc:
            raise DataError(str(exc)) from None
        return _from_grid(name, grid, header_row)
    if ext == ".xls":
        raise DataError("eski .xls biçimi desteklenmiyor: dosyayı Excel'de .xlsx olarak kaydedip tekrar dene")
    if ext in TEXT_EXT:
        text, enc = _read_text(p)
        delim = _sniff(text, ext)
        grid = [[c if c != "" else None for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]
        table = _from_grid(p.stem, grid, header_row)
        table.notes.insert(0, f"metin tablo: kodlama {enc}, ayraç {delim!r}")
        return table
    if ext == ".json":
        try:
            data = json.loads(p.read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as exc:
            raise DataError(f"geçersiz JSON: {exc}") from None
        if json_key:
            if not isinstance(data, dict) or json_key not in data:
                raise DataError(f"JSON'da '{json_key}' anahtarı yok")
            data = data[json_key]
        if isinstance(data, dict):
            raise DataError("JSON kökü bir nesne: kayıt listesini göstermek için --json-key <anahtar> ver")
        if not isinstance(data, list) or not all(isinstance(x, dict) for x in data):
            raise DataError("JSON bir nesne listesi olmalı: [{...}, {...}]")
        columns: list[str] = []
        for rec in data:
            for key in rec:
                if key not in columns:
                    columns.append(key)

        def scalar(v):
            return json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v
        return Table(p.stem, columns, [[scalar(rec.get(c)) for c in columns] for rec in data])
    raise UsageError(f"desteklenmeyen uzantı '{ext}': .xlsx .xlsm .csv .tsv .txt .json")


def check_output(inp, out, force: bool, allowed: set[str]) -> Path:
    out_p = Path(out)
    if out_p.suffix.lower() not in allowed:
        raise UsageError(f"çıktı uzantısı {out_p.suffix or '(yok)'} desteklenmiyor: {' '.join(sorted(allowed))}")
    for i in ([inp] if not isinstance(inp, (list, tuple)) else inp):
        if Path(i).resolve() == out_p.resolve():
            raise UsageError("girdi dosyasının üzerine yazılmaz: farklı bir çıktı yolu ver")
    if out_p.exists() and not force:
        raise UsageError(f"çıktı zaten var: {out_p} (ezmek için --force)")
    return out_p


def save_table(t: Table, out: Path, sheet_name: str | None = None, delimiter: str = ",") -> list[str]:
    out.parent.mkdir(parents=True, exist_ok=True)
    ext = out.suffix.lower()
    if ext == ".xlsx":
        return X.write_xlsx(out, [(sheet_name or t.name or "Sayfa1", [t.columns] + t.rows)])
    if ext in (".csv", ".tsv"):
        with out.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter="\t" if ext == ".tsv" else delimiter)
            writer.writerow(t.columns)
            writer.writerows([[show(v) for v in r] for r in t.rows])
        return []
    out.write_text(json.dumps([dict(zip(t.columns, r)) for r in t.rows], ensure_ascii=False, indent=2), encoding="utf-8")
    return []


def print_preview(t: Table, n: int = 5) -> None:
    if n <= 0 or not t.rows:
        return
    print(f"İlk {min(n, len(t.rows))} satır:")
    print("  " + " | ".join(t.columns))
    for r in t.rows[:n]:
        print("  " + " | ".join(str(show(v)) for v in r))


# ── profile ───────────────────────────────────────────────────────────────────────────────────────

def cmd_profile(a) -> int:
    p = Path(a.file)
    sheets = []
    if p.suffix.lower() in EXCEL_EXT and p.is_file():
        try:
            sheets = X.list_sheets(p)
        except X.XlsxError as exc:
            raise DataError(str(exc)) from None
    t = load_table(p, a.sheet, a.header_row, a.json_key)
    columns = []
    for j, name in enumerate(t.columns):
        values = [r[j] for r in t.rows]
        filled = [v for v in values if not blank(v)]
        kinds = sorted({kind(v) for v in filled})
        info = {"kolon": name,
                "tip": "boş" if not kinds else (kinds[0] if len(kinds) == 1 else "karışık(" + "+".join(kinds) + ")"),
                "dolu": len(filled), "boş": len(values) - len(filled), "farklı_değer": len({str(v) for v in filled})}
        if not a.no_samples:
            samples: list[str] = []
            for v in filled:
                if str(v) not in samples:
                    samples.append(str(v))
                if len(samples) == 3:
                    break
            info["örnek"] = samples
        nums = [to_number(v) for v in filled if kind(v) == "sayı"]
        if nums and not a.no_samples:  # min/max/toplam da değer sızdırır (ör. tek satırlık kimlik no)
            info.update({"min": min(nums), "max": max(nums), "toplam": round(sum(nums), 6)})
        columns.append(info)
    warnings = list(t.notes)
    empty_cols = [c["kolon"] for c in columns if c["dolu"] == 0]
    if t.columns and (len(empty_cols) * 2 > len(t.columns) or t.blank_headers * 2 > len(t.columns)):
        warnings.append(f"kolonların yarısından fazlası boş ya da başlıksız (boş {len(empty_cols)}, başlıksız "
                        f"{t.blank_headers} / {len(t.columns)}): başlık {a.header_row}. satırda olmayabilir "
                        "(ör. SAP BEx/ALV dışa aktarımı) → --header-row N ile yeniden profille")
    mixed = [c["kolon"] for c in columns if c["tip"].startswith("karışık")]
    if mixed:
        warnings.append(f"karışık tipli kolonlar (sayı ve metin bir arada): {', '.join(mixed)}")
    result = {"dosya": str(p), "sayfalar": sheets, "sayfa": t.name, "başlık_satırı": a.header_row,
              "veri_satırı": len(t.rows), "kolon_sayısı": len(t.columns), "kolonlar": columns, "uyarılar": warnings}
    if not a.no_samples and a.rows > 0:
        result["ilk_satırlar"] = [dict(zip(t.columns, [show(v) for v in r])) for r in t.rows[:a.rows]]
    if a.json:
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        return 0
    if sheets:
        print(f"Sayfalar: {', '.join(sheets)}")
    print(f"Sayfa: {t.name} · başlık satırı: {a.header_row} · veri satırı: {len(t.rows)} · kolon: {len(t.columns)}")
    for c in columns:
        extra = f" · min {c['min']} · max {c['max']} · toplam {c['toplam']}" if "min" in c else ""
        sample = f" · örnek: {', '.join(c['örnek'])}" if c.get("örnek") else ""
        print(f"  - {c['kolon']}: {c['tip']} · dolu {c['dolu']} · boş {c['boş']} · farklı {c['farklı_değer']}{extra}{sample}")
    for w in warnings:
        print(f"UYARI: {w}")
    if not a.no_samples:
        print_preview(t, a.rows)
    return 0


# ── convert ───────────────────────────────────────────────────────────────────────────────────────

def cmd_convert(a) -> int:
    t = load_table(a.input, a.sheet, a.header_row, a.json_key)
    out = check_output(a.input, a.output, a.force, {".xlsx", ".csv", ".tsv", ".json"})
    for w in t.notes + save_table(t, out, a.out_sheet, a.csv_delimiter):
        print(f"NOT: {w}")
    print(f"yazıldı: {out} · {len(t.rows)} satır · {len(t.columns)} kolon")
    return 0


# ── transform ─────────────────────────────────────────────────────────────────────────────────────

OPS = ("eq", "ne", "gt", "ge", "lt", "le", "contains", "startswith", "empty", "notempty")
AGGS = ("sum", "count", "mean", "min", "max", "first")


def _match(v, op: str, target: str) -> bool:
    if op == "empty":
        return blank(v)
    if op == "notempty":
        return not blank(v)
    if op in ("contains", "startswith"):
        s, tv = tr_lower("" if v is None else str(v)), tr_lower(target)
        return tv in s if op == "contains" else s.startswith(tv)
    x, y = to_number(v), to_number(target)
    if x is None or y is None:
        if op in ("gt", "ge", "lt", "le") and blank(v):
            return False
        x, y = ("" if v is None else str(v).strip()), target.strip()
    return {"eq": x == y, "ne": x != y, "gt": x > y, "ge": x >= y, "lt": x < y, "le": x <= y}[op]


def _sort_rows(rows: list[list], j: int, desc: bool) -> list[list]:
    filled = [r for r in rows if not blank(r[j])]
    empty = [r for r in rows if blank(r[j])]

    def key(r):
        n = to_number(r[j])
        return (0, n, "") if n is not None else (1, 0, tr_lower(str(r[j])))
    return sorted(filled, key=key, reverse=desc) + empty


def cmd_transform(a) -> int:
    t = load_table(a.input, a.sheet, a.header_row, a.json_key)
    out = check_output(a.input, a.output, a.force, {".xlsx", ".csv", ".tsv", ".json"})
    before = (len(t.rows), len(t.columns))
    rows = t.rows
    for expr in a.where or []:
        parts = expr.split("|", 2)
        if len(parts) < 2 or parts[1] not in OPS or (parts[1] not in ("empty", "notempty") and len(parts) != 3):
            raise UsageError(f"--where biçimi 'Kolon|op|değer' (op: {', '.join(OPS)}): {expr!r}")
        j = t.col(parts[0])
        rows = [r for r in rows if _match(r[j], parts[1], parts[2] if len(parts) == 3 else "")]
    if a.dedupe or a.dedupe_on:
        idx = [t.col(c) for c in a.dedupe_on] if a.dedupe_on else list(range(len(t.columns)))
        seen, unique = set(), []
        for r in rows:
            k = tuple(str(show(r[i])) for i in idx)
            if k not in seen:
                seen.add(k)
                unique.append(r)
        rows = unique
    columns = list(t.columns)
    if a.group_by:
        g_idx = [t.col(c) for c in a.group_by]
        aggs = []
        for spec in a.agg or []:
            m = re.match(r"^(.+?)=(\w+):(.+)$", spec)
            if not m or m.group(2) not in AGGS:
                raise UsageError(f"--agg biçimi 'YeniAd=fonksiyon:Kolon' (fonksiyon: {', '.join(AGGS)}): {spec!r}")
            aggs.append((m.group(1), m.group(2), t.col(m.group(3))))
        groups: dict[tuple, list[list]] = {}
        for r in rows:
            groups.setdefault(tuple(show(r[i]) for i in g_idx), []).append(r)
        new_rows = []
        for key, members in groups.items():
            out_row = list(key)
            for _name, fn, j in aggs:
                vals = [r[j] for r in members if not blank(r[j])]
                nums = [to_number(v) for v in vals if to_number(v) is not None]
                if fn == "count":
                    out_row.append(len(vals))
                elif fn == "first":
                    out_row.append(vals[0] if vals else None)
                elif fn in ("min", "max") and len(nums) != len(vals):
                    out_row.append((min if fn == "min" else max)(str(v) for v in vals) if vals else None)
                elif not nums:
                    out_row.append(None)
                else:
                    out_row.append({"sum": sum(nums), "mean": sum(nums) / len(nums), "min": min(nums), "max": max(nums)}[fn])
            new_rows.append(out_row)
        columns = list(a.group_by) + [name for name, _fn, _j in aggs]
        rows = new_rows
    elif a.agg:
        raise UsageError("--agg yalnız --group-by ile kullanılır")
    work = Table(t.name, columns, rows)
    for spec in reversed(a.sort or []):
        name, _, direction = spec.rpartition(":") if spec.lower().endswith((":desc", ":asc")) else (spec, "", "asc")
        rows = _sort_rows(rows, work.col(name), direction.lower() == "desc")
    for spec in a.rename or []:
        old, sep, new = spec.partition("=")
        if not sep or not new:
            raise UsageError(f"--rename biçimi 'Eski=Yeni': {spec!r}")
        columns[work.col(old)] = new
    work = Table(t.name, columns, rows)
    if a.keep:
        idx = [work.col(c) for c in a.keep]
    else:
        drop = {work.col(c) for c in (a.drop or [])}
        idx = [j for j in range(len(columns)) if j not in drop]
    result = Table(t.name, [columns[j] for j in idx], [[r[j] for j in idx] for r in rows])
    for w in t.notes + save_table(result, out, a.out_sheet, a.csv_delimiter):
        print(f"NOT: {w}")
    print(f"yazıldı: {out} · satır {before[0]} → {len(result.rows)} · kolon {before[1]} → {len(result.columns)}")
    print_preview(result, a.preview)
    return 0


# ── compare ───────────────────────────────────────────────────────────────────────────────────────

def _key(row: list, idx: list[int], mode: str) -> tuple:
    parts = []
    for j in idx:
        s = "" if row[j] is None else str(row[j])
        if mode in ("trim", "number"):
            s = s.strip()
        if mode == "number":
            n = to_number(s)
            if n is not None:
                s = str(int(n)) if float(n).is_integer() else str(n)
        parts.append(s)
    return tuple(parts)


def _same(x, y, strict: bool) -> bool:
    if strict:
        return str(show(x)) == str(show(y))
    if blank(x) and blank(y):
        return True
    nx, ny = to_number(x), to_number(y)
    if nx is not None and ny is not None:
        return abs(nx - ny) <= 1e-9 * max(1.0, abs(nx), abs(ny))
    return str(show(x)).strip() == str(show(y)).strip()


def cmd_compare(a) -> int:
    old = load_table(a.old, a.sheet_old, a.header_row, a.json_key)
    new = load_table(a.new, a.sheet_new, a.header_row, a.json_key)
    out = None
    if a.output and not a.summary_only:
        out = check_output([a.old, a.new], a.output, a.force, {".xlsx"})
    k_old = [old.col(k) for k in a.key]
    k_new = [new.col(k) for k in a.key]
    common = [c for c in old.columns if c in new.columns and c not in a.key]
    only_old = [c for c in old.columns if c not in new.columns]
    only_new = [c for c in new.columns if c not in old.columns]

    def index(t: Table, idx: list[int]):
        mapping, dups = {}, []
        for r in t.rows:
            k = _key(r, idx, a.key_mode)
            if k in mapping:
                dups.append(k)
            else:
                mapping[k] = r
        return mapping, dups
    m_old, d_old = index(old, k_old)
    m_new, d_new = index(new, k_new)
    added = [m_new[k] for k in m_new if k not in m_old]
    removed = [m_old[k] for k in m_old if k not in m_new]
    changes = []
    changed_rows = 0
    for k, r_old in m_old.items():
        r_new = m_new.get(k)
        if r_new is None:
            continue
        diff = [(c, r_old[old.col(c)], r_new[new.col(c)]) for c in common
                if not _same(r_old[old.col(c)], r_new[new.col(c)], a.strict)]
        if diff:
            changed_rows += 1
            changes.extend((k, c, x, y) for c, x, y in diff)
    same_rows = sum(1 for k in m_old if k in m_new) - changed_rows
    summary = {"satır_eski": len(old.rows), "satır_yeni": len(new.rows), "eklenen": len(added), "silinen": len(removed),
               "değişen_satır": changed_rows, "değişen_hücre": len(changes), "aynı_satır": same_rows,
               "tekrarlanan_anahtar_eski": len(d_old), "tekrarlanan_anahtar_yeni": len(d_new),
               "yalnız_eskide_kolon": only_old, "yalnız_yenide_kolon": only_new, "anahtar_modu": a.key_mode}
    for label, value in summary.items():
        print(f"{label}: {value}")
    if d_old or d_new:
        print("UYARI: anahtar tekil değil; tekrarlanan anahtarlarda yalnız İLK satır karşılaştırıldı → bileşik anahtar dene "
              "(--key A --key B)")
    if only_old or only_new:
        print("UYARI: kolon kümeleri farklı; yalnız ortak kolonlar karşılaştırıldı")
    if not a.summary_only:
        for k, c, x, y in changes[:a.show]:
            print(f"  değişti {'/'.join(k)} · {c}: {show(x)!r} → {show(y)!r}")
        if len(changes) > a.show > 0:
            print(f"  … {len(changes) - a.show} değişiklik daha")
    if out:
        sheets = [
            ("Ozet", [["Ölçü", "Değer"]] + [[k, ", ".join(v) if isinstance(v, list) else v] for k, v in summary.items()]),
            ("Eklenen", [new.columns] + added),
            ("Silinen", [old.columns] + removed),
            ("Degisen", [list(a.key) + ["Kolon", "Eski", "Yeni"]] + [list(k) + [c, x, y] for k, c, x, y in changes]),
        ]
        if d_old or d_new:
            sheets.append(("Tekrarlanan_Anahtar", [["Dosya"] + list(a.key)] + [["eski"] + list(k) for k in d_old]
                           + [["yeni"] + list(k) for k in d_new]))
        X.write_xlsx(out, sheets)
        print(f"fark raporu: {out} (satır verisi içerir — paylaşmadan önce içeriği kontrol et)")
    return 0


# ── images ────────────────────────────────────────────────────────────────────────────────────────

def _safe_name(s: str) -> str:
    return re.sub(r"[^\w.\-]+", "_", s).strip("_") or "resim"


def cmd_images(a) -> int:
    src = Path(a.file)
    if not src.is_file():
        raise UsageError(f"dosya yok: {src}")
    if src.suffix.lower() == ".xls":
        raise DataError("eski .xls biçimi desteklenmiyor: dosyayı .xlsx olarak kaydedip tekrar dene")
    out = Path(a.out) if a.out else src.with_name(src.stem + "_images")
    if out.exists() and any(out.iterdir()) and not a.force:
        raise UsageError(f"çıktı klasörü boş değil: {out} (--force ya da başka --out)")
    records = []
    with X.open_zip(src) as zf:
        names = set(zf.namelist())
        media = sorted(n for n in names if n.startswith("xl/media/"))
        placements = []
        for sheet_name, part in X.sheet_parts(zf):
            for target, typ in X.part_rels(zf, part).values():
                if not typ.endswith("/drawing") or target not in names:
                    continue
                drawing_rels = X.part_rels(zf, target)
                for anchor in list(ET.fromstring(zf.read(target))):
                    frm = anchor.find(XDR + "from")
                    if frm is not None:
                        cell = f"{X.col_letters(int(frm.findtext(XDR + 'col', '0')))}{int(frm.findtext(XDR + 'row', '0')) + 1}"
                    else:
                        cell = "konumsuz"
                    for blip in anchor.iter(A_NS + "blip"):
                        tgt = drawing_rels.get(blip.get(X.R + "embed"), (None, ""))[0]
                        if tgt and tgt in names:
                            placements.append((sheet_name, cell, tgt))
        if not media:
            print("gömülü resim yok (xl/media boş). Grafik, koşullu biçim ve hücre rengi resim değildir.")
            return 0
        out.mkdir(parents=True, exist_ok=True)
        written: dict[str, str] = {}
        counters: dict[str, int] = {}

        def write(data: bytes, base: str, ext: str) -> str:
            digest = hashlib.sha256(data).hexdigest()
            if digest not in written:
                counters[base] = counters.get(base, 0) + 1
                fname = _safe_name(f"{base}_{counters[base]}") + ext
                (out / fname).write_bytes(data)
                written[digest] = fname
            return written[digest]
        for sheet_name, cell, tgt in placements:
            fname = write(zf.read(tgt), f"{sheet_name}_{cell}", Path(tgt).suffix.lower())
            records.append({"sayfa": sheet_name, "hücre": cell, "dosya": fname, "kaynak": tgt})
        placed = {t for _s, _c, t in placements}
        for m in media:
            if m not in placed:
                fname = write(zf.read(m), "konumsuz_" + Path(m).stem, Path(m).suffix.lower())
                records.append({"sayfa": None, "hücre": None, "dosya": fname, "kaynak": m})
    if a.json:
        print(json.dumps({"klasör": str(out), "resimler": records, "benzersiz_dosya": len(written)}, ensure_ascii=False, indent=2))
    else:
        print(f"klasör: {out} · benzersiz resim: {len(written)} · yerleşim: {len(records)}")
        for r in records:
            where = f"{r['sayfa']}!{r['hücre']}" if r["sayfa"] else "hücreye bağlı değil (üst/alt bilgi ya da grup)"
            print(f"  - {r['dosya']} ← {where}")
    return 0


# ── report (openpyxl) ─────────────────────────────────────────────────────────────────────────────

def cmd_report(a) -> int:
    try:
        import openpyxl
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.properties import PageSetupProperties
    except ImportError:
        raise DepMissing("openpyxl") from None
    t = load_table(a.input, a.sheet, a.header_row, a.json_key)
    out = check_output(a.input, a.output, a.force, {".xlsx"})
    for color in (a.header_color, a.alt_row_color):
        if not re.fullmatch(r"#?[0-9A-Fa-f]{6}", color):
            raise UsageError(f"renk 6 haneli hex olmalı (ör. 1F4E79): {color!r}")
    formats = {}
    for spec in a.number_format or []:
        col, sep, fmt = spec.partition("=")
        if not sep or not fmt:
            raise UsageError(f"--number-format biçimi 'Kolon=biçim': {spec!r}")
        t.col(col)
        formats[col] = fmt
    if a.chart and (not a.chart_x or not a.chart_y):
        raise UsageError("--chart için --chart-x ve en az bir --chart-y gerekli")
    numeric = {t.col(c) for c in list(formats) + list(a.chart_y or [])}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = X.safe_sheet_name(a.out_sheet or t.name or "Rapor", set())[0]
    first = 1
    if a.title:
        ws.cell(row=1, column=1, value=a.title).font = Font(bold=True, size=14)
        if len(t.columns) > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(t.columns))
        first = 3
    head_fill = PatternFill("solid", fgColor=a.header_color.lstrip("#"))
    alt_fill = PatternFill("solid", fgColor=a.alt_row_color.lstrip("#"))
    for j, name in enumerate(t.columns, start=1):
        cell = ws.cell(row=first, column=j, value=name)
        cell.fill, cell.font = head_fill, Font(bold=True, color="FFFFFF")
    for i, row in enumerate(t.rows):
        for j, value in enumerate(row):
            if j in numeric and to_number(value) is not None:
                value = to_number(value)
            cell = ws.cell(row=first + 1 + i, column=j + 1, value=value)
            if t.columns[j] in formats:
                cell.number_format = formats[t.columns[j]]
            if i % 2 == 1:
                cell.fill = alt_fill
    last = first + len(t.rows)
    ws.freeze_panes = ws.cell(row=first + 1, column=1)
    if t.rows and t.columns:
        ws.auto_filter.ref = f"A{first}:{get_column_letter(len(t.columns))}{last}"
    for j, name in enumerate(t.columns, start=1):
        longest = max([len(str(name))] + [len(str(show(r[j - 1]))) for r in t.rows[:2000]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 2, 8), 60)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if a.chart and t.rows:
        chart = BarChart() if a.chart == "bar" else LineChart()
        chart.title = a.title or None
        for y in a.chart_y:
            chart.add_data(Reference(ws, min_col=t.col(y) + 1, min_row=first, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=t.col(a.chart_x) + 1, min_row=first + 1, max_row=last))
        chart.width, chart.height = 18, 9
        ws.add_chart(chart, f"{get_column_letter(len(t.columns) + 2)}{first}")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"rapor: {out} · {len(t.rows)} satır · biçimli kolon: {', '.join(formats) or '-'} · grafik: {a.chart or '-'}")
    return 0


# ── komut satırı ──────────────────────────────────────────────────────────────────────────────────

def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description="Excel/CSV/JSON tablo işleri (office-excel skill'i)")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def table_args(p, sheet_flag="--sheet"):
        p.add_argument(sheet_flag, help="xlsx sayfa adı ya da 1'den başlayan sıra no (varsayılan: ilk sayfa)")
        p.add_argument("--header-row", type=int, default=1, help="başlık satırı no (varsayılan 1)")
        p.add_argument("--json-key", help="JSON kökü nesneyse kayıt listesinin anahtarı")

    p = sub.add_parser("profile", help="dosyayı profille")
    p.add_argument("file")
    table_args(p)
    p.add_argument("--rows", type=int, default=5, help="gösterilecek örnek satır (varsayılan 5)")
    p.add_argument("--no-samples", action="store_true", help="değer gösterme (hassas veri): yalnız sayılar ve tipler")
    p.add_argument("--json", action="store_true")
    p.set_defaults(func=cmd_profile)

    def output_args(p):
        p.add_argument("--out-sheet", help="xlsx çıktıda sayfa adı")
        p.add_argument("--csv-delimiter", default=",", help="csv çıktı ayracı (varsayılan ,)")
        p.add_argument("--force", action="store_true", help="var olan çıktıyı ez")

    p = sub.add_parser("convert", help="biçim dönüştür")
    p.add_argument("input")
    p.add_argument("output")
    table_args(p)
    output_args(p)
    p.set_defaults(func=cmd_convert)

    p = sub.add_parser("transform", help="süz / grupla / sırala / kolon seç")
    p.add_argument("input")
    p.add_argument("output")
    table_args(p)
    p.add_argument("--where", action="append", help="'Kolon|op|değer' (tekrarlanabilir, VE): op " + " ".join(OPS))
    p.add_argument("--dedupe", action="store_true", help="tüm kolonlarda aynı satırları at")
    p.add_argument("--dedupe-on", action="append", help="bu kolon(lar)a göre tekrarları at")
    p.add_argument("--group-by", action="append", help="gruplama kolonu (tekrarlanabilir)")
    p.add_argument("--agg", action="append", help="'YeniAd=fonksiyon:Kolon' fonksiyon: " + " ".join(AGGS))
    p.add_argument("--sort", action="append", help="'Kolon' ya da 'Kolon:desc' (ilk verilen birincil)")
    p.add_argument("--rename", action="append", help="'Eski=Yeni'")
    p.add_argument("--keep", action="append", help="yalnız bu kolonlar (bu sırayla)")
    p.add_argument("--drop", action="append", help="bu kolonları at")
    p.add_argument("--preview", type=int, default=5, help="çıktıdan gösterilecek satır (0 = gösterme)")
    output_args(p)
    p.set_defaults(func=cmd_transform)

    p = sub.add_parser("compare", help="iki tabloyu anahtarla karşılaştır")
    p.add_argument("old")
    p.add_argument("new")
    p.add_argument("--key", action="append", required=True, help="anahtar kolon (bileşik anahtar için tekrarla)")
    p.add_argument("--sheet-old")
    p.add_argument("--sheet-new")
    p.add_argument("--header-row", type=int, default=1)
    p.add_argument("--json-key")
    p.add_argument("--key-mode", choices=("exact", "trim", "number"), default="trim",
                   help="anahtar normalleştirme: number = baştaki sıfırları yok say (000010 = 10)")
    p.add_argument("--strict", action="store_true", help="değerleri metin olarak birebir karşılaştır (1 ≠ 1.0)")
    p.add_argument("--output", help="fark raporu .xlsx")
    p.add_argument("--summary-only", action="store_true", help="yalnız sayılar; satır verisi basılmaz/yazılmaz")
    p.add_argument("--show", type=int, default=5, help="ekrana basılacak değişiklik sayısı")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_compare)

    p = sub.add_parser("images", help="gömülü resimleri çıkar")
    p.add_argument("file")
    p.add_argument("--out", help="çıktı klasörü (varsayılan: <dosya>_images)")
    p.add_argument("--json", action="store_true")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_images)

    p = sub.add_parser("report", help="biçimli rapor (openpyxl gerekir)")
    p.add_argument("input")
    p.add_argument("output")
    table_args(p)
    p.add_argument("--title")
    p.add_argument("--number-format", action="append", help="'Kolon=#,##0.00' (tekrarlanabilir)")
    p.add_argument("--header-color", default="1F4E79")
    p.add_argument("--alt-row-color", default="EEF3F8")
    p.add_argument("--chart", choices=("bar", "line"))
    p.add_argument("--chart-x", help="grafik kategori kolonu")
    p.add_argument("--chart-y", action="append", help="grafik değer kolonu (tekrarlanabilir)")
    p.add_argument("--out-sheet")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_report)
    return ap


def main(argv=None) -> int:
    ap = build_parser()
    try:
        args = ap.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code in (0, None) else 3
    try:
        return args.func(args) or 0
    except UsageError as exc:
        print(f"HATA (kullanım): {exc}", file=sys.stderr)
        return 3
    except DataError as exc:
        print(f"HATA (veri): {exc}", file=sys.stderr)
        return 1
    except DepMissing as exc:
        print(f"EKSİK BAĞIMLILIK: bu alt komut '{exc}' paketini ister. Kurmak kullanıcının kararıdır: "
              f"python -m pip install --user {exc}", file=sys.stderr)
        return 4
    except PermissionError as exc:
        print(f"HATA: dosyaya yazılamadı ({exc}) — dosya Excel'de açık olabilir", file=sys.stderr)
        return 1
    except (zipfile.BadZipFile, ET.ParseError) as exc:
        print(f"HATA (veri): dosya çözümlenemedi ({type(exc).__name__}: {exc})", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())

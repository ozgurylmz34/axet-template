#!/usr/bin/env python3
"""md_to_docx.py — Markdown (md_lite alt kümesi) → düzenlenebilir Word .docx.

İsteğe bağlı bağımlılık: python-docx (`python -m pip install --user python-docx` — kurmak kullanıcının kararı).
Başlıklar, liste, tablo ("Table Grid"), kod bloğu (Consolas), alıntı, yatay çizgi, sayfa sonu, yerel görsel
(sayfa genişliğine sığdırılır, alt metin altyazı olur). Bağlantılar altı çizili metin olur (tıklanabilir değil).

Kullanım:
  python md_to_docx.py --input rapor.md --output rapor.docx [--title "Başlık"] [--redact-pii [--pii-mode genis]] [--force]
Çıkış: 0 başarı · 1 dönüştürme hatası · 3 kullanım hatası · 4 python-docx yok.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import md_lite  # noqa: E402
import pii_redact  # noqa: E402

DOCX_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}


def _style(doc, *names):
    for name in names:
        try:
            doc.styles[name]
            return name
        except KeyError:
            continue
    return None


def build(blocks: list[dict], out: Path, title: str | None, base_dir: Path) -> dict:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()
    doc.core_properties.author = ""
    if title:
        doc.core_properties.title = title
        doc.add_heading(title, 0)
    stats = {"başlık": 0, "tablo": 0, "görsel": 0, "uyarı": []}
    align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}

    def add_runs(paragraph, text, bold=False):
        for r in md_lite.inline_runs(text):
            run = paragraph.add_run(r["text"])
            run.bold = r["bold"] or bold or None
            run.italic = r["italic"] or None
            if r["code"]:
                run.font.name = "Consolas"
            if r["href"]:
                run.underline = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    for b in blocks:
        kind = b["type"]
        if kind == "heading":
            add_runs(doc.add_heading("", min(b["level"], 9)), b["text"])
            stats["başlık"] += 1
        elif kind == "paragraph":
            add_runs(doc.add_paragraph(), b["text"])
        elif kind == "list":
            base = "List Number" if b["ordered"] else "List Bullet"
            for level, text in b["items"]:
                name = base if level == 0 else f"{base} {min(level + 1, 3)}"
                add_runs(doc.add_paragraph(style=_style(doc, name, base)), text)
        elif kind == "table":
            table = doc.add_table(rows=1 + len(b["rows"]), cols=len(b["header"]))
            style = _style(doc, "Table Grid")
            if style:
                table.style = style
            for i, row in enumerate([b["header"]] + b["rows"]):
                for j, cell_text in enumerate(row):
                    paragraph = table.cell(i, j).paragraphs[0]
                    paragraph.alignment = align_map[b["align"][j]]
                    add_runs(paragraph, cell_text, bold=(i == 0))
            stats["tablo"] += 1
        elif kind == "code":
            run = doc.add_paragraph().add_run(b["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif kind == "quote":
            add_runs(doc.add_paragraph(style=_style(doc, "Intense Quote", "Quote")), b["text"])
        elif kind == "hr":
            p_pr = doc.add_paragraph()._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(key), value)
            border.append(bottom)
            p_pr.append(border)
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "image":
            path, src = md_lite.resolve_image(b["src"], base_dir)
            if path is None or not path.is_file() or path.suffix.lower() not in DOCX_IMAGE_EXT:
                stats["uyarı"].append(f"görsel eklenmedi (yok, uzak ya da desteklenmeyen biçim): {src}")
                doc.add_paragraph(f"[görsel eklenmedi: {src}]")
                continue
            shape = doc.add_picture(str(path))
            max_width = Inches(6)
            if shape.width > max_width:
                ratio = max_width / shape.width
                shape.width, shape.height = int(max_width), int(shape.height * ratio)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            alt = md_lite.restore(b["alt"])
            if alt:
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.add_run(alt).italic = True
            stats["görsel"] += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return stats


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Markdown → Word .docx (python-docx)")
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--title", help="belgenin başına Title stilinde başlık")
    ap.add_argument("--redact-pii", action="store_true")
    ap.add_argument("--pii-mode", choices=pii_redact.MODES, default="akilli")
    ap.add_argument("--force", action="store_true")
    try:
        args = ap.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code in (0, None) else 3
    src, out = Path(args.input), Path(args.output)
    if not src.is_file():
        print(f"HATA (kullanım): girdi yok: {src}", file=sys.stderr)
        return 3
    if out.suffix.lower() != ".docx":
        print("HATA (kullanım): çıktı .docx olmalı", file=sys.stderr)
        return 3
    if out.exists() and not args.force:
        print(f"HATA (kullanım): çıktı zaten var: {out} (ezmek için --force)", file=sys.stderr)
        return 3
    try:
        import docx  # noqa: F401
    except ImportError:
        print("EKSİK BAĞIMLILIK: python-docx gerekli. Kurmak kullanıcının kararıdır: "
              "python -m pip install --user python-docx", file=sys.stderr)
        return 4
    text = src.read_text(encoding="utf-8-sig")
    if args.redact_pii:
        text, counts = pii_redact.redact_text(text, args.pii_mode)
        print("maskeleme: " + " · ".join(f"{k}={v}" for k, v in counts.items()))
    try:
        stats = build(md_lite.parse(text), out, args.title, src.parent)
    except PermissionError as exc:
        print(f"HATA: yazılamadı ({exc}) — dosya Word'de açık olabilir", file=sys.stderr)
        return 1
    for w in stats.pop("uyarı"):
        print(f"UYARI: {w}")
    print(f"yazıldı: {out} · " + " · ".join(f"{k} {v}" for k, v in stats.items()))
    return 0


if __name__ == "__main__":
    sys.exit(main())

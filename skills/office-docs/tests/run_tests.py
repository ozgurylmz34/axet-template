#!/usr/bin/env python3
"""office-docs çevrimdışı testleri (stdlib unittest).

    python skills/office-docs/tests/run_tests.py

Örnek girdi: tests/samples/ornek.md (görsel test sırasında üretilir). python-docx yoksa DOCX testleri, Edge/Chrome
yoksa PDF testi SKIP olur. Çıkış: 0 tümü geçti · 1 en az bir başarısız.
"""
from __future__ import annotations

import importlib.util
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
SAMPLE = HERE / "samples" / "ornek.md"
sys.path.insert(0, str(SCRIPTS))
import md_lite  # noqa: E402
import md_to_pdf  # noqa: E402
import pii_redact  # noqa: E402

HAS_DOCX = importlib.util.find_spec("docx") is not None
BROWSER = md_to_pdf.find_browser()


def make_png(w: int = 40, h: int = 20) -> bytes:
    raw = b"".join(b"\x00" + b"\x1f\x4e\x79" * w for _ in range(h))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def run(script: str, *args, code: str | None = None):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
    cmd = [sys.executable, "-c", code] if code else [sys.executable, str(SCRIPTS / script)]
    p = subprocess.run(cmd + [str(a) for a in args], capture_output=True, text=True, encoding="utf-8",
                       stdin=subprocess.DEVNULL, env=env, timeout=300)
    return p.returncode, p.stdout, p.stderr


class Tmp(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.md = self.d / "ornek.md"
        shutil.copy(SAMPLE, self.md)
        (self.d / "ekran.png").write_bytes(make_png())

    def tearDown(self):
        self._tmp.cleanup()


class ParseTests(unittest.TestCase):
    def setUp(self):
        self.blocks = md_lite.parse(SAMPLE.read_text(encoding="utf-8"))
        self.types = [b["type"] for b in self.blocks]

    def test_block_sequence(self):
        self.assertEqual(self.types, ["heading", "paragraph", "heading", "list", "list", "table", "quote", "code",
                                      "image", "pagebreak", "heading", "hr", "paragraph"])

    def test_nested_list_levels(self):
        self.assertEqual(self.blocks[3]["items"], [(0, "Toplam malzeme: 3"), (1, "Çelik vida"),
                                                   (1, "Şaft ğ ü ş ı İ ö ç"), (0, "Açık kalem: 0")])
        self.assertTrue(self.blocks[4]["ordered"])

    def test_table_align_and_escaped_pipe(self):
        table = self.blocks[5]
        self.assertEqual(table["align"], ["left", "center", "right"])
        self.assertEqual(md_lite.plain_text(table["rows"][0][1]), "Vida | M8")

    def test_inline_runs(self):
        runs = md_lite.inline_runs(self.blocks[1]["text"])
        styled = {(r["text"], r["bold"], r["italic"], r["code"], r["href"]) for r in runs}
        self.assertIn(("örnek", True, False, False, None), styled)
        self.assertIn(("İstanbul", False, True, False, None), styled)
        self.assertIn(("ZCA000_CLC", False, False, True, None), styled)
        self.assertIn(("abapGit", False, False, False, "https://docs.abapgit.org/"), styled)

    def test_escape_and_snake_case(self):
        self.assertEqual(md_lite.plain_text(self.blocks[-1]["text"]), "Son satır *yıldız* kaçışı ve snake_case_ad.")

    def test_html_embeds_image_and_escapes(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "ekran.png").write_bytes(make_png())
            doc, warnings = md_lite.to_html(self.blocks + [{"type": "paragraph", "text": "<script>x</script>"}], "T", tmp)
        self.assertEqual(warnings, [])
        self.assertIn('src="data:image/png;base64,', doc)
        self.assertIn("&lt;script&gt;", doc)
        self.assertIn("<ul><li>Toplam malzeme: 3<ul><li>Çelik vida</li><li>Şaft ğ ü ş ı İ ö ç</li></ul></li>", doc)

    def test_missing_image_warns(self):
        _doc, warnings = md_lite.to_html([{"type": "image", "alt": "", "src": "yok.png"}], "T", ".")
        self.assertEqual(len(warnings), 1)


class RedactTests(unittest.TestCase):
    TEXT = SAMPLE.read_text(encoding="utf-8")

    def test_tckn_checksum(self):
        self.assertTrue(pii_redact.tckn_valid("10000000146"))
        self.assertFalse(pii_redact.tckn_valid("10000000147"))
        self.assertFalse(pii_redact.tckn_valid("05321234567"))

    def test_smart_mode_keeps_sap_document_numbers(self):
        masked, counts = pii_redact.redact_text(self.TEXT)
        self.assertNotIn("10000000146", masked)
        self.assertNotIn("1234567890", masked)
        self.assertNotIn("TR33 0006", masked)
        self.assertIn("0080001234", masked)
        self.assertIn("000010", masked)
        self.assertEqual(counts, {"iban": 1, "tckn": 1, "vkn": 1, "10_11_hane": 0})

    def test_broad_mode_masks_all_10_11_digit_runs(self):
        masked, counts = pii_redact.redact_text(self.TEXT, "genis")
        self.assertNotIn("0080001234", masked)
        self.assertEqual(counts["10_11_hane"], 3)
        self.assertIn("000010", masked)  # 6 hane korunur

    def test_invalid_11_digits_untouched_in_smart_mode(self):
        self.assertEqual(pii_redact.redact_text("no 12345678901")[0], "no 12345678901")


class DocxTests(Tmp):
    @unittest.skipUnless(HAS_DOCX, "python-docx yok")
    def test_build_docx(self):
        import docx
        out = self.d / "rapor.docx"
        rc, stdout, err = run("md_to_docx.py", "--input", self.md, "--output", out, "--title", "Stok", "--redact-pii")
        self.assertEqual(rc, 0, err)
        self.assertIn("tablo 1", stdout)
        self.assertIn("görsel 1", stdout)
        document = docx.Document(str(out))
        texts = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Şaft ğ ü ş ı İ ö ç", texts)
        self.assertNotIn("10000000146", texts)
        self.assertIn("0080001234", texts)
        self.assertEqual(document.tables[0].cell(1, 1).text, "Vida | M8")
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(rc, 0)
        self.assertEqual(run("md_to_docx.py", "--input", self.md, "--output", out)[0], 3)

    def test_docx_dependency_missing_exit_4(self):
        code = ("import sys, runpy; sys.modules['docx'] = None; "
                f"sys.argv = ['md_to_docx.py', '--input', {str(self.md)!r}, '--output', {str(self.d / 'x.docx')!r}]; "
                f"runpy.run_path({str(SCRIPTS / 'md_to_docx.py')!r}, run_name='__main__')")
        rc, _, err = run("", code=code)
        self.assertEqual(rc, 4)
        self.assertIn("python-docx", err)


class PdfTests(Tmp):
    @unittest.skipUnless(BROWSER, "Edge/Chrome bulunamadı")
    def test_build_pdf(self):
        out = self.d / "rapor.pdf"
        html = self.d / "rapor.html"
        rc, stdout, err = run("md_to_pdf.py", "--input", self.md, "--output", out, "--keep-html", html, "--redact-pii")
        self.assertEqual(rc, 0, err + stdout)
        data = out.read_bytes()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 5000)
        self.assertNotIn("10000000146", html.read_text(encoding="utf-8"))
        self.assertIn("sayfa: 2", stdout)  # <!-- pagebreak --> ikinci sayfayı açar

    def test_bad_browser_path_is_usage_error(self):
        rc, _, err = run("md_to_pdf.py", "--input", self.md, "--output", self.d / "x.pdf", "--browser", self.d / "yok.exe")
        self.assertEqual(rc, 3)


if __name__ == "__main__":
    unittest.main(verbosity=2)

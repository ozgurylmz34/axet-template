#!/usr/bin/env python3
"""build_pptx.py — düzenlenebilir PowerPoint (.pptx): JSON spesifikasyondan, Markdown desteden ya da tablo verisinden.

İsteğe bağlı bağımlılık: python-pptx (`python -m pip install --user python-pptx` — kurmak kullanıcının kararı).
Spesifikasyon üretimi (`--md`, `--table`) ve `--dump-spec` python-pptx olmadan da çalışır.

Slayt tipleri (16:9, tüm kutular gerçek metin kutusu/tablo — PowerPoint'te düzenlenebilir):
  {"type": "title",   "title": "...", "subtitle": "..."}
  {"type": "bullets", "title": "...", "bullets": ["madde", ["alt madde", 1]]}
  {"type": "table",   "title": "...", "headers": [...], "rows": [[...]]}      (uzun tablo "devam" slaytlarına bölünür)
  {"type": "image",   "title": "...", "image": "ekran.png", "caption": "..."} (görsel spesifikasyon dosyasına görelidir)

Kullanım:
  python build_pptx.py --spec deste.json --output deste.pptx
  python build_pptx.py --md deste.md --output deste.pptx            (--- slayt ayırır; # başlık; ## alt başlık; - madde)
  python build_pptx.py --table veri.csv --title "Özet" --group-col Bölge --output deste.pptx
  ... [--redact-pii [--pii-mode genis]] [--max-rows 12] [--dump-spec spec.json] [--force]
Çıkış: 0 başarı · 1 hata · 3 kullanım hatası · 4 python-pptx yok.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pii_redact  # noqa: E402

SLIDE_TYPES = ("title", "bullets", "table", "image")
_LIST = re.compile(r"^(\s*)(?:[-*+]|\d+[.)])\s+(.*)$")
_IMAGE = re.compile(r'^\s*!\[([^\]]*)\]\(\s*([^)\s]+)(?:\s+"[^"]*")?\s*\)\s*$')
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_INLINE_MARKS = re.compile(r"\*\*|__|`")


class UsageError(Exception):
    pass


def _clean(text: str) -> str:
    return _INLINE_MARKS.sub("", text).strip()


def _cells(line: str) -> list[str]:
    s = line.strip().strip("|")
    return [_clean(c.replace("\\|", "|")) for c in re.split(r"(?<!\\)\|", s)]


def md_to_spec(text: str) -> dict:
    chunks, current = [], []
    for line in text.replace("\r\n", "\n").split("\n"):
        if re.match(r"^\s*---\s*$", line):
            chunks.append(current)
            current = []
        else:
            current.append(line)
    chunks.append(current)
    slides = []
    for chunk in chunks:
        title = subtitle = None
        bullets, tables, images = [], [], []
        i = 0
        while i < len(chunk):
            line = chunk[i]
            if not line.strip():
                i += 1
                continue
            if line.startswith("# "):
                title = _clean(line[2:])
            elif line.startswith("## "):
                subtitle = _clean(line[3:])
            elif "|" in line and i + 1 < len(chunk) and _TABLE_SEP.match(chunk[i + 1]):
                headers, rows, i = _cells(line), [], i + 2
                while i < len(chunk) and "|" in chunk[i] and chunk[i].strip():
                    rows.append((_cells(chunk[i]) + [""] * len(headers))[:len(headers)])
                    i += 1
                tables.append((headers, rows))
                continue
            elif _IMAGE.match(line):
                m = _IMAGE.match(line)
                images.append((m.group(2), _clean(m.group(1))))
            elif _LIST.match(line):
                m = _LIST.match(line)
                bullets.append((len(m.group(1).expandtabs(4)), _clean(m.group(2))))
            else:
                bullets.append((-1, _clean(line)))
            i += 1
        if not any((title, subtitle, bullets, tables, images)):
            continue
        indents = sorted({ind for ind, _t in bullets if ind >= 0})
        items = [t if ind <= (indents[0] if indents else 0) else [t, min(indents.index(ind), 4)] for ind, t in bullets]
        if bullets:
            slides.append({"type": "bullets", "title": title or "", "bullets": items})
        for headers, rows in tables:
            slides.append({"type": "table", "title": title or "", "headers": headers, "rows": rows})
        for src, caption in images:
            slides.append({"type": "image", "title": title or "", "image": src, "caption": caption})
        if not (bullets or tables or images):
            slides.append({"type": "title", "title": title or "", "subtitle": subtitle or ""})
    return {"slides": slides}


def load_table(path: Path) -> tuple[list[str], list[list]]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        raise UsageError("Excel girdisini önce office-excel ile CSV'ye çevir: office_excel.py convert veri.xlsx veri.csv")
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1254")
    if ext == ".json":
        data = json.loads(text)
        if not isinstance(data, list) or not all(isinstance(x, dict) for x in data):
            raise UsageError("JSON bir nesne listesi olmalı: [{...}]")
        columns: list[str] = []
        for rec in data:
            columns += [k for k in rec if k not in columns]
        return columns, [[rec.get(c, "") for c in columns] for rec in data]
    if ext not in (".csv", ".tsv", ".txt"):
        raise UsageError(f"desteklenmeyen tablo uzantısı: {ext}")
    try:
        delim = "\t" if ext == ".tsv" else csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        delim = ","
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]
    if not rows:
        raise UsageError("tablo boş")
    return rows[0], rows[1:]


def table_to_spec(columns: list[str], rows: list[list], title: str, group_col: str | None = None) -> dict:
    slides = [{"type": "title", "title": title, "subtitle": f"{len(rows)} satır · {len(columns)} kolon"}]
    if group_col:
        if group_col not in columns:
            raise UsageError(f"kolon yok: {group_col!r} · kolonlar: {', '.join(columns)}")
        g = columns.index(group_col)
        groups: dict[str, list[list]] = {}
        for r in rows:
            groups.setdefault(str(r[g]), []).append(r)
        slides.append({"type": "table", "title": "Özet", "headers": [group_col, "Satır"],
                       "rows": [[k, str(len(v))] for k, v in groups.items()]})
        rest = [c for c in columns if c != group_col]
        for key, members in groups.items():
            slides.append({"type": "table", "title": f"{group_col}: {key}", "headers": rest,
                           "rows": [[str(r[j]) for j, c in enumerate(columns) if c != group_col] for r in members]})
    else:
        slides.append({"type": "table", "title": title, "headers": list(columns), "rows": [[str(v) for v in r] for r in rows]})
    return {"title": title, "slides": slides}


def validate(spec: dict) -> list[str]:
    errors = []
    slides = spec.get("slides") if isinstance(spec, dict) else None
    if not isinstance(slides, list) or not slides:
        return ["spesifikasyonda 'slides' listesi yok ya da boş"]
    for n, s in enumerate(slides, start=1):
        kind = s.get("type") if isinstance(s, dict) else None
        if kind not in SLIDE_TYPES:
            errors.append(f"slayt {n}: type {kind!r} geçersiz ({', '.join(SLIDE_TYPES)})")
        elif kind == "bullets" and not isinstance(s.get("bullets"), list):
            errors.append(f"slayt {n}: 'bullets' listesi yok")
        elif kind == "table" and (not isinstance(s.get("headers"), list) or not isinstance(s.get("rows"), list)):
            errors.append(f"slayt {n}: 'headers' ve 'rows' listeleri gerekli")
        elif kind == "image" and not s.get("image"):
            errors.append(f"slayt {n}: 'image' yolu yok")
    return errors


def redact_spec(obj, mode: str, totals: dict):
    if isinstance(obj, str):
        masked, counts = pii_redact.redact_text(obj, mode)
        for k, v in counts.items():
            totals[k] = totals.get(k, 0) + v
        return masked
    if isinstance(obj, list):
        return [redact_spec(x, mode, totals) for x in obj]
    if isinstance(obj, dict):
        return {k: (v if k in ("type", "image") else redact_spec(v, mode, totals)) for k, v in obj.items()}
    return obj


def build(spec: dict, out: Path, base_dir: Path, max_rows: int) -> dict:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    accent = RGBColor(0x1F, 0x4E, 0x79)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    warnings: list[str] = []

    def textbox(slide, left, top, width, height, text, size, bold=False, color=None, align=None):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        p = frame.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold = Pt(size), bold
        if color is not None:
            p.font.color.rgb = color
        if align is not None:
            p.alignment = align
        return frame

    def header(slide, title):
        textbox(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9), title, 28, True, accent)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

    for s in spec["slides"]:
        kind = s["type"]
        if kind == "title":
            slide = prs.slides.add_slide(blank)
            textbox(slide, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.4), s.get("title", ""), 40, True, accent, PP_ALIGN.CENTER)
            if s.get("subtitle"):
                textbox(slide, Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.9), s["subtitle"], 20, align=PP_ALIGN.CENTER)
        elif kind == "bullets":
            slide = prs.slides.add_slide(blank)
            header(slide, s.get("title", ""))
            frame = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(12), Inches(5.6)).text_frame
            frame.word_wrap = True
            for n, item in enumerate(s["bullets"]):
                text, level = (item, 0) if isinstance(item, str) else (str(item[0]), int(item[1]) if len(item) > 1 else 0)
                p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                p.text = ("• " if level == 0 else "– ") + text
                p.level = max(0, min(level, 4))
                p.font.size = Pt(max(20 - 2 * p.level, 12))
        elif kind == "table":
            headers, rows = [str(h) for h in s["headers"]], s["rows"]
            pages = [rows[i:i + max_rows] for i in range(0, len(rows), max_rows)] or [[]]
            for n, page in enumerate(pages, start=1):
                slide = prs.slides.add_slide(blank)
                suffix = f" (devam {n}/{len(pages)})" if len(pages) > 1 else ""
                header(slide, s.get("title", "") + suffix)
                shape = slide.shapes.add_table(len(page) + 1, len(headers), Inches(0.5), Inches(1.5), Inches(12.3),
                                               Inches(0.4) * (len(page) + 1))
                table = shape.table
                for j, h in enumerate(headers):
                    table.cell(0, j).text = h
                for i, row in enumerate(page, start=1):
                    for j in range(len(headers)):
                        table.cell(i, j).text = str(row[j]) if j < len(row) and row[j] is not None else ""
                for i in range(len(page) + 1):
                    for j in range(len(headers)):
                        for p in table.cell(i, j).text_frame.paragraphs:
                            p.font.size = Pt(12)
                            p.font.bold = i == 0
        elif kind == "image":
            slide = prs.slides.add_slide(blank)
            header(slide, s.get("title", ""))
            path = Path(s["image"])
            path = path if path.is_absolute() else base_dir / path
            if not path.is_file():
                warnings.append(f"görsel yok: {s['image']}")
                textbox(slide, Inches(0.7), Inches(3), Inches(12), Inches(1), f"[görsel bulunamadı: {s['image']}]", 18)
                continue
            pic = slide.shapes.add_picture(str(path), Inches(0.5), Inches(1.45))
            box_w, box_h = Inches(12.3), Inches(5.0 if s.get("caption") else 5.7)
            scale = min(box_w / pic.width, box_h / pic.height)
            pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
            pic.left = Emu(int(Inches(0.5) + (box_w - pic.width) / 2))
            if s.get("caption"):
                textbox(slide, Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.6), s["caption"], 14, align=PP_ALIGN.CENTER)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return {"slayt": len(prs.slides), "uyarılar": warnings}


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Düzenlenebilir PowerPoint üret (python-pptx)")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--spec", help="JSON slayt spesifikasyonu")
    src.add_argument("--md", help="Markdown deste (--- slayt ayırır)")
    src.add_argument("--table", help="csv/tsv/json tablo → özet deste")
    ap.add_argument("--output", help=".pptx çıktı (--dump-spec tek başına verilmişse gerekmez)")
    ap.add_argument("--title", help="deste başlığı (--table için zorunlu değil, varsayılan dosya adı)")
    ap.add_argument("--group-col", help="--table: her grup için ayrı slayt")
    ap.add_argument("--max-rows", type=int, default=12, help="tablo slaydı başına satır (varsayılan 12)")
    ap.add_argument("--redact-pii", action="store_true")
    ap.add_argument("--pii-mode", choices=pii_redact.MODES, default="akilli")
    ap.add_argument("--dump-spec", help="üretilen spesifikasyonu bu JSON dosyasına da yaz")
    ap.add_argument("--force", action="store_true")
    try:
        args = ap.parse_args(argv)
    except SystemExit as exc:
        return 0 if exc.code in (0, None) else 3
    try:
        if args.spec:
            base = Path(args.spec).resolve().parent
            spec = json.loads(Path(args.spec).read_text(encoding="utf-8-sig"))
        elif args.md:
            base = Path(args.md).resolve().parent
            spec = md_to_spec(Path(args.md).read_text(encoding="utf-8-sig"))
        else:
            base = Path(args.table).resolve().parent
            columns, rows = load_table(Path(args.table))
            spec = table_to_spec(columns, rows, args.title or Path(args.table).stem, args.group_col)
        if args.title and isinstance(spec, dict) and not args.table:
            spec["title"] = args.title
        errors = validate(spec)
        if errors:
            raise UsageError("; ".join(errors))
        if args.max_rows < 1:
            raise UsageError("--max-rows en az 1")
        if args.redact_pii:
            totals: dict = {}
            spec = redact_spec(spec, args.pii_mode, totals)
            print("maskeleme: " + " · ".join(f"{k}={v}" for k, v in totals.items()))
        if args.dump_spec:
            Path(args.dump_spec).write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"spesifikasyon: {args.dump_spec} · {len(spec['slides'])} slayt tanımı")
        if not args.output:
            if args.dump_spec:
                return 0
            raise UsageError("--output gerekli")
        out = Path(args.output)
        if out.suffix.lower() != ".pptx":
            raise UsageError("çıktı .pptx olmalı")
        if out.exists() and not args.force:
            raise UsageError(f"çıktı zaten var: {out} (ezmek için --force)")
    except (UsageError, FileNotFoundError, json.JSONDecodeError) as exc:
        print(f"HATA (kullanım): {exc}", file=sys.stderr)
        return 3
    try:
        import pptx  # noqa: F401
    except ImportError:
        print("EKSİK BAĞIMLILIK: python-pptx gerekli. Kurmak kullanıcının kararıdır: "
              "python -m pip install --user python-pptx (spesifikasyon için --dump-spec bağımlılıksız çalışır)",
              file=sys.stderr)
        return 4
    try:
        result = build(spec, out, base, args.max_rows)
    except PermissionError as exc:
        print(f"HATA: yazılamadı ({exc}) — dosya PowerPoint'te açık olabilir", file=sys.stderr)
        return 1
    for w in result["uyarılar"]:
        print(f"UYARI: {w}")
    print(f"yazıldı: {out} · {result['slayt']} slayt")
    return 0


if __name__ == "__main__":
    sys.exit(main())

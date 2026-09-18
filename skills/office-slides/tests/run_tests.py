#!/usr/bin/env python3
"""office-slides çevrimdışı testleri (stdlib unittest).

    python skills/office-slides/tests/run_tests.py

Spesifikasyon üretimi her ortamda test edilir; .pptx üretimi python-pptx kuruluysa (yoksa SKIP).
Çıkış: 0 tümü geçti · 1 en az bir başarısız.
"""
from __future__ import annotations

import importlib.util
import json
import os
import struct
import subprocess
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
import build_pptx as B  # noqa: E402

HAS_PPTX = importlib.util.find_spec("pptx") is not None
DECK_MD = """# Çeyrek Değerlendirmesi
## Yönetim özeti

---

# Öne çıkanlar
- Ciro **%12** arttı
  - Ege bölgesi öncü
- Müşteri VKN: 1234567890 raporda görünmemeli

---

# Skor tablosu
| Ölçü | Değer |
|---|---|
| Başarılı | 22 |
| Hatalı \\| açık | 0 |

---

# Ekran
![SAP ekranı](ekran.png)
"""


def make_png(w: int = 64, h: int = 32) -> bytes:
    raw = b"".join(b"\x00" + b"\x10\x80\x30" * w for _ in range(h))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def run(*args, code: str | None = None):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
    cmd = [sys.executable, "-c", code] if code else [sys.executable, str(SCRIPTS / "build_pptx.py")]
    p = subprocess.run(cmd + [str(a) for a in args], capture_output=True, text=True, encoding="utf-8",
                       stdin=subprocess.DEVNULL, env=env)
    return p.returncode, p.stdout, p.stderr


class SpecTests(unittest.TestCase):
    def test_md_to_spec(self):
        spec = B.md_to_spec(DECK_MD)
        self.assertEqual([s["type"] for s in spec["slides"]], ["title", "bullets", "table", "image"])
        self.assertEqual(spec["slides"][0], {"type": "title", "title": "Çeyrek Değerlendirmesi", "subtitle": "Yönetim özeti"})
        self.assertEqual(spec["slides"][1]["bullets"], ["Ciro %12 arttı", ["Ege bölgesi öncü", 1],
                                                       "Müşteri VKN: 1234567890 raporda görünmemeli"])
        self.assertEqual(spec["slides"][2]["rows"], [["Başarılı", "22"], ["Hatalı | açık", "0"]])
        self.assertEqual(spec["slides"][3]["caption"], "SAP ekranı")
        self.assertEqual(B.validate(spec), [])

    def test_table_to_spec_grouped(self):
        spec = B.table_to_spec(["Bölge", "Ürün", "Adet"], [["Ege", "Vida", "3"], ["Marmara", "Pul", "1"], ["Ege", "Pul", "2"]],
                               "Satış", "Bölge")
        self.assertEqual([s["title"] for s in spec["slides"]], ["Satış", "Özet", "Bölge: Ege", "Bölge: Marmara"])
        self.assertEqual(spec["slides"][1]["rows"], [["Ege", "2"], ["Marmara", "1"]])
        self.assertEqual(spec["slides"][2]["headers"], ["Ürün", "Adet"])

    def test_validate_errors(self):
        self.assertTrue(B.validate({"slides": [{"type": "chart"}]}))
        self.assertTrue(B.validate({}))

    def test_redact_spec_keeps_paths(self):
        totals: dict = {}
        spec = B.redact_spec({"slides": [{"type": "image", "image": "10000000146.png", "caption": "TCKN 10000000146"}]},
                             "akilli", totals)
        self.assertEqual(spec["slides"][0]["image"], "10000000146.png")
        self.assertNotIn("10000000146", spec["slides"][0]["caption"])
        self.assertEqual(totals["tckn"], 1)

    def test_redactor_copy_identical_to_office_docs(self):
        other = HERE.parents[1] / "office-docs" / "scripts" / "pii_redact.py"
        self.assertEqual((SCRIPTS / "pii_redact.py").read_bytes(), other.read_bytes(),
                         "pii_redact.py iki skill'de farklılaştı: birini diğerine eşitle")

    def test_dump_spec_without_pptx_and_xlsx_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "v.csv").write_text("Bölge;Adet\nEge;3\n", encoding="utf-8")
            code = ("import sys, runpy; sys.modules['pptx'] = None; "
                    f"sys.argv = ['b', '--table', {str(d / 'v.csv')!r}, '--dump-spec', {str(d / 's.json')!r}]; "
                    f"runpy.run_path({str(SCRIPTS / 'build_pptx.py')!r}, run_name='__main__')")
            rc, _, err = run(code=code)
            self.assertEqual(rc, 0, err)
            self.assertEqual(len(json.loads((d / "s.json").read_text(encoding="utf-8"))["slides"]), 2)
            (d / "v.xlsx").write_bytes(b"PK")
            rc, _, err = run("--table", d / "v.xlsx", "--output", d / "x.pptx")
            self.assertEqual(rc, 3)
            self.assertIn("office-excel", err)

    def test_missing_pptx_exit_4(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "d.md").write_text(DECK_MD, encoding="utf-8")
            code = ("import sys, runpy; sys.modules['pptx'] = None; "
                    f"sys.argv = ['b', '--md', {str(d / 'd.md')!r}, '--output', {str(d / 'd.pptx')!r}]; "
                    f"runpy.run_path({str(SCRIPTS / 'build_pptx.py')!r}, run_name='__main__')")
            rc, _, err = run(code=code)
            self.assertEqual(rc, 4)
            self.assertIn("python-pptx", err)


@unittest.skipUnless(HAS_PPTX, "python-pptx yok")
class BuildTests(unittest.TestCase):
    def test_build_from_md_with_image_and_redaction(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "d.md").write_text(DECK_MD, encoding="utf-8")
            (d / "ekran.png").write_bytes(make_png())
            out = d / "d.pptx"
            rc, stdout, err = run("--md", d / "d.md", "--output", out, "--redact-pii")
            self.assertEqual(rc, 0, err)
            self.assertIn("4 slayt", stdout)
            prs = Presentation(str(out))
            texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
            joined = "\n".join(texts)
            self.assertIn("Çeyrek Değerlendirmesi", joined)
            self.assertIn("– Ege bölgesi öncü", joined)
            self.assertNotIn("1234567890", joined)
            tables = [sh.table for s in prs.slides for sh in s.shapes if sh.has_table]
            self.assertEqual(tables[0].cell(2, 0).text, "Hatalı | açık")
            pictures = [sh for sh in prs.slides[3].shapes if sh.shape_type == 13]
            self.assertEqual(len(pictures), 1)
            self.assertLessEqual(pictures[0].width, int(12.3 * 914400) + 1)
            self.assertEqual(run("--md", d / "d.md", "--output", out)[0], 3)

    def test_long_table_splits(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            (d / "v.csv").write_text("No,Ad\n" + "".join(f"{i},Satır {i}\n" for i in range(25)), encoding="utf-8")
            out = d / "t.pptx"
            rc, _, err = run("--table", d / "v.csv", "--title", "Uzun", "--max-rows", "10", "--output", out)
            self.assertEqual(rc, 0, err)
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 4)  # başlık + 3 tablo sayfası
            titles = [sh.text_frame.text for sh in prs.slides[3].shapes if sh.has_text_frame]
            self.assertIn("Uzun (devam 3/3)", titles)


if __name__ == "__main__":
    unittest.main(verbosity=2)
